"""Update planning CV with explicit city/country on roles and degrees."""
from pathlib import Path

from docx import Document

CV_PATH = Path(
    r"c:\Users\gamal\OneDrive - Qatar University\CV\Gamal_Farag_Planning_CV_Portfolio_2026.docx"
)
REPO_PDF = Path(__file__).resolve().parents[1] / "assets" / "Gamal_Farag_CV.pdf"

REPLACEMENTS = {
    "AL FAISAL HOLDING  |  Planning Engineer | Owner Representative Office\tMar 2024 – Present": (
        "AL FAISAL HOLDING  |  Doha, Qatar  |  Planning Engineer  |  Owner Representative Office  |  Mar 2024 – Present"
    ),
    "AL FAISAL HOLDING  |  Planning Engineer | Owner Representative Office\tMar 2024 - Present": (
        "AL FAISAL HOLDING  |  Doha, Qatar  |  Planning Engineer  |  Owner Representative Office  |  Mar 2024 – Present"
    ),
    "Doha, Qatar  |  Private Royal Developments / Four Palaces Programme (QAR 200M+)": (
        "Private Royal Developments / Four Palaces Programme (QAR 200M+)"
    ),
    "MANNAI CORPORATION / GULF LABORATORIES  |  Planning Engineer (Civil)\tSep 2021 – Mar 2024": (
        "MANNAI CORPORATION / GULF LABORATORIES  |  Doha, Qatar  |  Planning Engineer (Civil)  |  Sep 2021 – Mar 2024"
    ),
    "MANNAI CORPORATION / GULF LABORATORIES  |  Planning Engineer (Civil)\tSep 2021 - Mar 2024": (
        "MANNAI CORPORATION / GULF LABORATORIES  |  Doha, Qatar  |  Planning Engineer (Civil)  |  Sep 2021 – Mar 2024"
    ),
    "Doha, Qatar  |  ASHGHAL Deep Injection Wells Programme (QAR 150M)": (
        "ASHGHAL Deep Injection Wells Programme (QAR 150M)"
    ),
    "CEG INTERNATIONAL  |  Planning Engineer (Civil)\tNov 2020 – Aug 2021": (
        "CEG INTERNATIONAL  |  Doha, Qatar  |  Planning Engineer (Civil)  |  Nov 2020 – Aug 2021"
    ),
    "CEG INTERNATIONAL  |  Planning Engineer (Civil)\tNov 2020 - Aug 2021": (
        "CEG INTERNATIONAL  |  Doha, Qatar  |  Planning Engineer (Civil)  |  Nov 2020 – Aug 2021"
    ),
    "Doha, Qatar  |  ASHGHAL Inner Doha Sewer Programme (QAR 500M)": (
        "ASHGHAL Foul Sewer Network For Inner Doha (QAR 500M)"
    ),
    "M.Sc. Engineering Management  |  Qatar University, 2024": (
        "Master's Degree (M.Sc.), Engineering Management  |  Qatar University  |  Doha, Qatar  |  2024"
    ),
    "B.Sc. Civil Engineering  |  Ain Shams University, 2020": (
        "Bachelor's Degree (B.Sc.), Civil Engineering  |  Ain Shams University, Faculty of Engineering  |  Cairo, Egypt  |  2020"
    ),
}


def replace_in_paragraph(paragraph, old: str, new: str) -> bool:
    if old not in paragraph.text:
        return False
    if paragraph.text.strip() == old.strip() or old in paragraph.text:
        paragraph.text = paragraph.text.replace(old, new)
        return True
    return False


def main():
    doc = Document(str(CV_PATH))
    for paragraph in doc.paragraphs:
        for old, new in REPLACEMENTS.items():
            replace_in_paragraph(paragraph, old, new)
        # GPA line may be separate part of same paragraph
    # Profile line — clarify master's location
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Planning Engineer (Civil) with nearly six years"):
            paragraph.text = paragraph.text.replace(
                "M.Sc. Engineering Management, Qatar University.",
                "Master's degree (M.Sc.) in Engineering Management, Qatar University, Doha, Qatar.",
            )

    doc.save(str(CV_PATH))
    print(f"Updated {CV_PATH}")

    try:
        from docx2pdf import convert

        convert(str(CV_PATH), str(REPO_PDF))
        print(f"Exported PDF to {REPO_PDF}")
    except Exception as exc:
        print(f"PDF export skipped ({exc}). Re-export CV to PDF manually from Word.")


if __name__ == "__main__":
    main()
